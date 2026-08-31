import os
import hashlib
import aiofiles
from datetime import datetime
from typing import List, Optional, Dict, Any
from uuid import UUID
import logging
import asyncio
from functools import partial

import chromadb
import trafilatura
import httpx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyMuPDFLoader, TextLoader
from langchain_ollama import OllamaEmbeddings
from langchain_core.documents import Document as LCDocument

from core.config import settings

logger = logging.getLogger(__name__)


class IndexingService:

    def __init__(self):
        self.chroma_client = chromadb.HttpClient(
            host=settings.chroma_host,
            port=settings.chroma_port
        )
        self.collection = self.chroma_client.get_or_create_collection(
            name=settings.chroma_collection,
            metadata={"hnsw:space": "cosine"}
        )
        self.embedder = OllamaEmbeddings(
            model="bge-m3:latest",
            base_url=settings.ollama_base_url,
        )
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
            separators=["\n\n", "\n", ".", "!", "?", " ", ""]
        )
        logger.info(f"IndexingService ready — bge-m3 @ {settings.ollama_base_url}")

    def compute_file_hash(self, file_path: str) -> str:
        md5 = hashlib.md5()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(8192), b""):
                md5.update(chunk)
        return md5.hexdigest()

    async def index_file(self, doc_id, file_path, filename, category, service_tag=None):
        try:
            text = await self._load_file(file_path, filename)
            return await self._index_text(doc_id=doc_id, text=text, filename=filename,
                category=category, service_tag=service_tag, source="file")
        except Exception as e:
            logger.error(f"Error indexing file {filename}: {e}")
            raise

    async def index_url(self, doc_id, url, title, category, service_tag=None):
        try:
            async with httpx.AsyncClient(timeout=30.0) as client:
                response = await client.get(url, follow_redirects=True)
                response.raise_for_status()
            text = trafilatura.extract(response.text, include_tables=True,
                include_links=False, output_format="markdown", favor_precision=True)
            if not text:
                raise ValueError(f"Could not extract content from {url}")
            full_text = f"# {title}\n\nSource: {url}\n\n{text}"
            return await self._index_text(doc_id=doc_id, text=full_text, filename=title,
                category=category, service_tag=service_tag, source="url", source_url=url)
        except Exception as e:
            logger.error(f"Error indexing URL {url}: {e}")
            raise

    async def delete_document(self, doc_id):
        try:
            results = self.collection.get(where={"doc_id": str(doc_id)})
            if results["ids"]:
                self.collection.delete(ids=results["ids"])
        except Exception as e:
            logger.error(f"Error deleting document {doc_id}: {e}")
            raise

    async def reindex_document(self, doc_id, file_path=None, url=None, filename="", category="other", service_tag=None):
        await self.delete_document(doc_id)
        if url:
            return await self.index_url(doc_id, url, filename, category, service_tag)
        elif file_path:
            return await self.index_file(doc_id, file_path, filename, category, service_tag)
        else:
            raise ValueError("Either file_path or url must be provided")

    def search(self, query, n_results=5, category_filter=None, service_filter=None):
        try:
            # embed_query is sync — ok for search (fast single query)
            query_embedding = self.embedder.embed_query(query)
            where = {}
            conditions = []
            if category_filter:
                conditions.append({"category": {"$eq": category_filter}})
            if service_filter:
                conditions.append({"service_tag": {"$eq": service_filter}})
            if len(conditions) == 1:
                where = conditions[0]
            elif len(conditions) > 1:
                where = {"$and": conditions}

            results = self.collection.query(
                query_embeddings=[query_embedding],
                n_results=n_results,
                where=where if where else None,
                include=["documents", "metadatas", "distances"],
            )

            if not results["ids"][0]:
                return []

            chunks = []
            for i, chunk_id in enumerate(results["ids"][0]):
                distance = results["distances"][0][i]
                relevance_score = round((1 - distance) * 100, 1)
                metadata = results["metadatas"][0][i]
                chunks.append({
                    "chunk_id": chunk_id,
                    "content": results["documents"][0][i],
                    "metadata": metadata,
                    "relevance_score": relevance_score,
                    "doc_id": metadata.get("doc_id", ""),
                    "filename": metadata.get("filename", ""),
                    "category": metadata.get("category", ""),
                    "page": metadata.get("page", 0),
                    "source_url": metadata.get("source_url", ""),
                    "chunk_index": metadata.get("chunk_index", 0),
                })
            return sorted(chunks, key=lambda x: x["relevance_score"], reverse=True)
        except Exception as e:
            logger.error(f"Search error: {e}")
            return []

    async def _load_file(self, file_path, filename):
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            # Run sync loader in thread to avoid blocking
            loop = asyncio.get_event_loop()
            pages = await loop.run_in_executor(None, lambda: PyMuPDFLoader(file_path).load())
            return "\n\n".join([p.page_content for p in pages])
        elif ext in [".txt", ".md", ".markdown"]:
            async with aiofiles.open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return await f.read()
        elif ext in [".docx", ".doc"]:
            from docx import Document as DocxDocument
            loop = asyncio.get_event_loop()
            def load_docx():
                doc = DocxDocument(file_path)
                return "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            return await loop.run_in_executor(None, load_docx)
        else:
            async with aiofiles.open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return await f.read()

    async def _index_text(self, doc_id, text, filename, category, service_tag, source, source_url=None):
        logger.info(f"Indexing {filename} — {len(text)} chars")

        lc_doc = LCDocument(page_content=text, metadata={"source": filename})
        chunks = self.splitter.split_documents([lc_doc])

        if not chunks:
            raise ValueError("No content could be extracted from document")

        logger.info(f"Split into {len(chunks)} chunks")

        texts = [chunk.page_content for chunk in chunks]
        ids = [f"{doc_id}_chunk_{i}" for i in range(len(chunks))]
        metadatas = [
            {
                "doc_id": str(doc_id),
                "filename": filename,
                "category": category,
                "service_tag": service_tag or "",
                "source": source,
                "source_url": source_url or "",
                "chunk_index": i,
                "page": chunk.metadata.get("page", 0),
                "indexed_at": datetime.utcnow().isoformat(),
            }
            for i, chunk in enumerate(chunks)
        ]

        # Run blocking embed_documents in a thread pool
        logger.info(f"Generating embeddings for {len(chunks)} chunks via bge-m3 (in thread)...")
        loop = asyncio.get_event_loop()
        embeddings = await loop.run_in_executor(
            None,
            partial(self.embedder.embed_documents, texts)
        )
        logger.info(f"Embeddings done for {len(chunks)} chunks")

        # Insert in batches
        batch_size = 100
        for i in range(0, len(ids), batch_size):
            self.collection.add(
                ids=ids[i:i + batch_size],
                embeddings=embeddings[i:i + batch_size],
                documents=texts[i:i + batch_size],
                metadatas=metadatas[i:i + batch_size],
            )
            logger.info(f"Inserted batch {i//batch_size + 1}/{(len(ids)+batch_size-1)//batch_size}")

        logger.info(f"Indexed {len(chunks)} chunks for {filename}")
        return len(chunks)


_indexing_service: Optional[IndexingService] = None


def get_indexing_service() -> IndexingService:
    global _indexing_service
    if _indexing_service is None:
        _indexing_service = IndexingService()
    return _indexing_service


indexing_service = None